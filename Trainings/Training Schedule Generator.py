import os
import calendar
import tkinter as tk
from datetime import datetime, timedelta
from tkinter import messagebox

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ---------------------------
# Helpers
# ---------------------------

def center_window(win, width=450, height=350):
    win.update_idletasks()
    sw = win.winfo_screenwidth()
    sh = win.winfo_screenheight()

    x = int((sw - width) / 2)
    y = int((sh - height) / 2)

    win.geometry(
        f"{width}x{height}+{x}+{y}"
    )


def format_date(date):
    return date.strftime(
        "%A, %B %d, %Y"
    ).replace(
        " 0",
        " "
    )


# ---------------------------
# Schedule Configuration
# ---------------------------

MILESTONES = [
    {
        "task": "Send Initial Email",
        "week": 0,
        "description": "Send the initial email to CELT-CDQ asking them to fill out the training schedule for the upcoming semester"
    },
    {
        "task": "Reminder Email",
        "week": 3,
        "description": "Send a reminder to CELT-CDQ members who have not yet completed the required information."
    },
    {
        "task": "Due Date",
        "week": 4,
        "description": "Deadline for CELT-CDQ members to submit their training information."
    },
    {
        "task": "Make Trainings",
        "week": 5,
        "description": "Create the teams meetings for the trainings based on the information submitted by CELT-CDQ members."
    },
    {
        "task": "Send Email",
        "week": 6,
        "description": "Send CELT central (Linds, Lisa, Krista) an email with their completed training information and any next steps."
    },
    {
        "task": "Due Date for Schedule",
        "week": 7,
        "description": "Final deadline for the training schedule to be completed and ready for the upcoming term."
    },
]


# ---------------------------
# Schedule Generation
# ---------------------------

def generate_schedule(schedule_due_date):
    schedule = []

    for milestone in MILESTONES:

        weeks_before_due = (
            7 - milestone["week"]
        )

        date = (
            schedule_due_date
            - timedelta(weeks=weeks_before_due)
        )

        schedule.append({
            "task": milestone["task"],
            "week": milestone["week"],
            "description": milestone["description"],
            "date": date
        })

    return schedule


# ---------------------------
# Schedule File
# ---------------------------

def create_schedule_file(
    semester,
    schedule
):

    script_folder = os.path.dirname(
        os.path.abspath(__file__)
    )

    folder_path = os.path.join(
        script_folder,
        semester
    )

    os.makedirs(
        folder_path,
        exist_ok=True
    )

    file_path = os.path.join(
        folder_path,
        "schedule.txt"
    )

    first_date = schedule[0]["date"]
    final_date = schedule[-1]["date"]

    with open(
        file_path,
        "w",
        encoding="utf-8"
    ) as file:

        # ---------------------------
        # Header
        # ---------------------------

        file.write("=" * 45 + "\n")
        file.write("           TRAINING SCHEDULE\n")
        file.write(
            f"           {semester.upper()}\n"
        )
        file.write("=" * 45 + "\n")
        file.write("\n")

        file.write(
            f"Schedule Due Date: "
            f"{format_date(final_date)}\n"
        )

        file.write("\n")

        # ---------------------------
        # At-a-Glance
        # ---------------------------

        file.write("-" * 45 + "\n")
        file.write("AT-A-GLANCE TIMELINE\n")
        file.write("-" * 45 + "\n")
        file.write("\n")

        file.write(
            f"{'Week':<7}"
            f"{'Date':<27}"
            f"Task\n"
        )

        file.write(
            f"{'-' * 5:<7}"
            f"{'-' * 24:<27}"
            f"{'-' * 30}\n"
        )

        for item in schedule:

            date_text = item["date"].strftime(
                "%A, %B %d"
            ).replace(
                " 0",
                " "
            )

            file.write(
                f"{item['week']:<7}"
                f"{date_text:<27}"
                f"{item['task']}\n"
            )

        file.write("\n")

        # ---------------------------
        # Detailed Timeline
        # ---------------------------

        file.write("-" * 45 + "\n")
        file.write("DETAILED TIMELINE\n")
        file.write("-" * 45 + "\n")
        file.write("\n")

        for item in schedule:

            file.write(
                f"Week {item['week']}\n"
            )

            file.write(
                f"{item['task']}\n"
            )

            file.write(
                f"{format_date(item['date'])}\n"
            )

            file.write(
                f"{item['description']}\n"
            )

            file.write("\n")

        # ---------------------------
        # Summary
        # ---------------------------

        file.write("-" * 45 + "\n")
        file.write("SUMMARY\n")
        file.write("-" * 45 + "\n")
        file.write("\n")

        file.write("First Action:\n")
        file.write(
            f"{format_date(first_date)}\n"
        )

        file.write("\n")

        file.write("Final Due Date:\n")
        file.write(
            f"{format_date(final_date)}\n"
        )

        file.write("\n")

        file.write("Timeline:\n")
        file.write("7 weeks\n")

    return file_path


# ---------------------------
# Word Hyperlink
# ---------------------------

def add_hyperlink(
    paragraph,
    text,
    url
):

    part = paragraph.part

    relationship_id = part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True
    )

    hyperlink = OxmlElement(
        "w:hyperlink"
    )

    hyperlink.set(
        qn("r:id"),
        relationship_id
    )

    run = OxmlElement(
        "w:r"
    )

    run_properties = OxmlElement(
        "w:rPr"
    )

    color = OxmlElement(
        "w:color"
    )

    color.set(
        qn("w:val"),
        "0563C1"
    )

    underline = OxmlElement(
        "w:u"
    )

    underline.set(
        qn("w:val"),
        "single"
    )

    run_properties.append(
        color
    )

    run_properties.append(
        underline
    )

    run.append(
        run_properties
    )

    text_element = OxmlElement(
        "w:t"
    )

    text_element.text = text

    run.append(
        text_element
    )

    hyperlink.append(
        run
    )

    paragraph._p.append(
        hyperlink
    )


# ---------------------------
# Template Metadata
# ---------------------------

def parse_template(
    template_text
):

    metadata = {}
    body_lines = []

    metadata_fields = [
        "Recipients",
        "Send Date",
        "Subject"
    ]

    in_body = False

    for line in template_text.splitlines():

        stripped = line.strip()

        if not in_body:

            found_metadata = False

            for field in metadata_fields:

                prefix = f"{field}:"

                if stripped.startswith(
                    prefix
                ):

                    metadata[field] = (
                        stripped[
                            len(prefix):
                        ].strip()
                    )

                    found_metadata = True
                    break

            if found_metadata:
                continue

            if stripped == "":
                continue

            in_body = True

        if in_body:
            body_lines.append(line)

    return metadata, body_lines


# ---------------------------
# Email Template Replacement
# ---------------------------

def replace_placeholders(
    text,
    semester,
    schedule
):

    dates = {
        item["week"]: item["date"]
        for item in schedule
    }

    replacements = {
        "[Semester]": semester,
        "[Week 0 Date]": format_date(
            dates[0]
        ),
        "[Week 3 Date]": format_date(
            dates[3]
        ),
        "[Week 6 Date]": format_date(
            dates[6]
        ),
        "[Due Date]": format_date(
            dates[4]
        )
    }

    for placeholder, value in replacements.items():

        text = text.replace(
            placeholder,
            value
        )

    return text


# ---------------------------
# Create Email DOCX
# ---------------------------

def create_email_docx(
    template_path,
    output_path,
    semester,
    schedule,
    spreadsheet_link
):

    with open(
        template_path,
        "r",
        encoding="utf-8"
    ) as file:

        template_text = file.read()

    metadata, body_lines = parse_template(
        template_text
    )

    required_metadata = [
        "Recipients",
        "Send Date",
        "Subject"
    ]

    for field in required_metadata:

        if field not in metadata:

            raise ValueError(
                f"Missing '{field}:' "
                f"in {os.path.basename(template_path)}"
            )

    recipients = replace_placeholders(
        metadata["Recipients"],
        semester,
        schedule
    )

    send_date = replace_placeholders(
        metadata["Send Date"],
        semester,
        schedule
    )

    subject = replace_placeholders(
        metadata["Subject"],
        semester,
        schedule
    )

    document = Document()

    # ---------------------------
    # Email Metadata
    # ---------------------------

    metadata_paragraphs = [
        ("To:", recipients),
        ("Send Date:", send_date),
        ("Subject:", subject)
    ]

    for label, value in metadata_paragraphs:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            f"{label} "
        )

        run.bold = True

        paragraph.add_run(
            value
        )

    document.add_paragraph()

    # ---------------------------
    # Email Body
    # ---------------------------

    for line in body_lines:

        paragraph = document.add_paragraph()

        placeholder = (
            "[Training Spreadsheet Link]"
        )

        line = replace_placeholders(
            line,
            semester,
            schedule
        )

        if placeholder in line:

            parts = line.split(
                placeholder
            )

            for index, part in enumerate(parts):

                if part:
                    paragraph.add_run(
                        part
                    )

                if index < len(parts) - 1:

                    add_hyperlink(
                        paragraph,
                        f"{semester} Trainings",
                        spreadsheet_link
                    )

        else:

            paragraph.add_run(
                line
            )

    document.save(
        output_path
    )

    return output_path


# ---------------------------
# Email Files
# ---------------------------

def create_email_files(
    semester,
    schedule,
    spreadsheet_link
):

    script_folder = os.path.dirname(
        os.path.abspath(__file__)
    )

    templates_folder = os.path.join(
        script_folder,
        "email_templates"
    )

    output_folder = os.path.join(
        script_folder,
        semester
    )

    os.makedirs(
        output_folder,
        exist_ok=True
    )

    templates = [
        (
            "01_initial_email.txt",
            "01_initial_email.docx"
        ),
        (
            "02_follow_up_email.txt",
            "02_follow_up_email.docx"
        ),
        (
            "03_training_links_email.txt",
            "03_training_links_email.docx"
        )
    ]

    output_files = []

    for template_name, output_name in templates:

        template_path = os.path.join(
            templates_folder,
            template_name
        )

        output_path = os.path.join(
            output_folder,
            output_name
        )

        if not os.path.exists(
            template_path
        ):

            raise FileNotFoundError(
                f"Email template not found:\n\n"
                f"{template_path}"
            )

        create_email_docx(
            template_path,
            output_path,
            semester,
            schedule,
            spreadsheet_link
        )

        output_files.append(
            output_path
        )

    return output_files


# ---------------------------
# Calendar Popup
# ---------------------------

class CalendarPopup:

    def __init__(
        self,
        parent,
        callback,
        selected_date
    ):

        self.parent = parent
        self.callback = callback
        self.selected_date = selected_date

        self.year = selected_date.year
        self.month = selected_date.month

        self.window = tk.Toplevel(
            parent
        )

        self.window.title(
            "Select Date"
        )

        self.window.resizable(
            False,
            False
        )

        self.window.transient(
            parent
        )

        self.create_calendar()

        center_window(
            self.window,
            450,
            285
        )

        self.window.grab_set()

    def create_calendar(self):

        header = tk.Frame(
            self.window
        )

        header.pack(
            fill="x",
            padx=10,
            pady=10
        )

        tk.Button(
            header,
            text="<",
            width=2,
            command=self.previous_month
        ).pack(
            side="left"
        )

        self.month_label = tk.Label(
            header,
            font=("Arial", 12, "bold")
        )

        self.month_label.pack(
            side="left",
            expand=True
        )

        tk.Button(
            header,
            text=">",
            width=3,
            command=self.next_month
        ).pack(
            side="right"
        )

        self.calendar_frame = tk.Frame(
            self.window
        )

        self.calendar_frame.pack(
            padx=10,
            pady=5
        )

        self.update_calendar()

    def update_calendar(self):

        for widget in self.calendar_frame.winfo_children():
            widget.destroy()

        self.month_label.config(
            text=(
                f"{calendar.month_name[self.month]} "
                f"{self.year}"
            )
        )

        weekdays = [
            "Mon",
            "Tue",
            "Wed",
            "Thu",
            "Fri",
            "Sat",
            "Sun"
        ]

        for column, day in enumerate(
            weekdays
        ):

            tk.Label(
                self.calendar_frame,
                text=day,
                width=3,
                font=("Arial", 9, "bold")
            ).grid(
                row=0,
                column=column,
                padx=1,
                pady=3
            )

        month_days = calendar.monthcalendar(
            self.year,
            self.month
        )

        for row, week in enumerate(
            month_days,
            start=1
        ):

            for column, day in enumerate(
                week
            ):

                if day == 0:
                    continue

                date = datetime(
                    self.year,
                    self.month,
                    day
                ).date()

                button = tk.Button(
                    self.calendar_frame,
                    text=str(day),
                    width=2,
                    command=lambda d=date:
                        self.select_date(d)
                )

                if date == self.selected_date:

                    button.config(
                        relief="sunken"
                    )

                button.grid(
                    row=row,
                    column=column,
                    padx=1,
                    pady=1
                )

    def previous_month(self):

        if self.month == 1:

            self.month = 12
            self.year -= 1

        else:

            self.month -= 1

        self.update_calendar()

    def next_month(self):

        if self.month == 12:

            self.month = 1
            self.year += 1

        else:

            self.month += 1

        self.update_calendar()

    def select_date(self, date):

        self.callback(date)

        self.window.grab_release()
        self.window.destroy()


# ---------------------------
# GUI
# ---------------------------

def get_inputs():

    root = tk.Tk()

    root.title(
        "Training Schedule Generator"
    )

    center_window(
        root,
        500,
        375
    )

    result = {
        "semester": None,
        "due_date": None,
        "spreadsheet_link": None
    }

    # ---------------------------
    # Semester
    # ---------------------------

    tk.Label(
        root,
        text="Semester:"
    ).pack(
        pady=(20, 3)
    )

    semester_entry = tk.Entry(
        root,
        width=30
    )

    semester_entry.pack()

    # ---------------------------
    # Due Date
    # ---------------------------

    tk.Label(
        root,
        text="Due Date for Schedule:"
    ).pack(
        pady=(15, 3)
    )

    date_frame = tk.Frame(
        root
    )

    date_frame.pack()

    selected_date = {
        "value": datetime.now().date()
    }

    date_label = tk.Label(
        date_frame,
        text=selected_date["value"].strftime(
            "%m/%d/%Y"
        ),
        width=12,
        relief="sunken",
        anchor="w"
    )

    date_label.pack(
        side="left"
    )

    def update_date(date):

        selected_date["value"] = date

        date_label.config(
            text=date.strftime(
                "%m/%d/%Y"
            )
        )

    tk.Button(
        date_frame,
        text="Select Date",
        command=lambda: CalendarPopup(
            root,
            update_date,
            selected_date["value"]
        )
    ).pack(
        side="left",
        padx=(5, 0)
    )

    # ---------------------------
    # Spreadsheet Link
    # ---------------------------

    tk.Label(
        root,
        text="Training Spreadsheet Link:"
    ).pack(
        pady=(15, 3)
    )

    spreadsheet_entry = tk.Entry(
        root,
        width=55
    )

    spreadsheet_entry.pack()

    # ---------------------------
    # Submit
    # ---------------------------

    def submit():

        semester = (
            semester_entry
            .get()
            .strip()
        )

        spreadsheet_link = (
            spreadsheet_entry
            .get()
            .strip()
        )

        if not semester:

            messagebox.showerror(
                "Error",
                "Enter semester"
            )

            return

        if not spreadsheet_link:

            messagebox.showerror(
                "Error",
                "Enter the training spreadsheet link"
            )

            return

        result["semester"] = semester
        result["due_date"] = (
            selected_date["value"]
        )
        result["spreadsheet_link"] = (
            spreadsheet_link
        )

        root.destroy()

    tk.Button(
        root,
        text="Generate Schedule",
        command=submit
    ).pack(
        pady=25
    )

    root.mainloop()

    return result


# ---------------------------
# Run
# ---------------------------

inputs = get_inputs()

if not inputs["semester"]:
    raise SystemExit(0)

semester = inputs["semester"]
schedule_due_date = inputs["due_date"]
spreadsheet_link = inputs["spreadsheet_link"]


# ---------------------------
# Generate Schedule
# ---------------------------

schedule_due_datetime = datetime.combine(
    schedule_due_date,
    datetime.min.time()
)

schedule = generate_schedule(
    schedule_due_datetime
)


# ---------------------------
# Create Files
# ---------------------------

try:

    schedule_file_path = (
        create_schedule_file(
            semester,
            schedule
        )
    )

    email_files = create_email_files(
        semester,
        schedule,
        spreadsheet_link
    )

except Exception as error:

    root = tk.Tk()
    root.withdraw()

    messagebox.showerror(
        "Error",
        f"Something went wrong:\n\n{error}"
    )

    root.destroy()

    raise SystemExit(1)


# ---------------------------
# Done
# ---------------------------

root = tk.Tk()
root.withdraw()

messagebox.showinfo(
    "Done",
    f"Files saved to:\n"
    f"{os.path.dirname(schedule_file_path)}"
)

root.destroy()