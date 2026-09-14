import os
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Alignment
from acad_departments import get_academic_department_data
from instructions import get_instructions


# ---------------------------
# Helpers
# ---------------------------

def colnum_to_letters(n: int) -> str:
    result = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        result = chr(65 + remainder) + result
    return result


def center_window(win, width=400, height=250):
    win.update_idletasks()
    sw = win.winfo_screenwidth()
    sh = win.winfo_screenheight()
    x = int((sw - width) / 2)
    y = int((sh - height) / 2)
    win.geometry(f"{width}x{height}+{x}+{y}")


def apply_conditional_formatting(ws, col_idx, max_row, max_col):
    fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    font = Font(color="9C0006")

    for r in range(2, max_row + 1):
        formula = f'=${colnum_to_letters(col_idx)}{r}="No"'
        ws.conditional_formatting.add(
            f"A{r}:{colnum_to_letters(max_col)}{r}",
            FormulaRule(formula=[formula], fill=fill, font=font)
        )


def create_table(ws, max_row, max_col, name):
    name = "".join(c if c.isalnum() else "_" for c in name)
    ref = f"A1:{colnum_to_letters(max_col)}{max_row}"

    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)


def ensure_columns(df, cols):
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise ValueError(f"Missing columns: {missing}")


# ---------------------------
# Sorting
# ---------------------------

def excel_like_sort(df,
                    dept,
                    subj,
                    num,
                    section,
                    lname):

    df = df.copy()

    for c in [dept, subj, lname, section]:
        df[c] = df[c].fillna("").astype(str).str.strip()

    df[num] = df[num].fillna("").astype(str).str.strip()

    df["_dept"] = df[dept].str.casefold()
    df["_subj"] = df[subj].str.casefold()
    df["_num"] = pd.to_numeric(df[num], errors="coerce")
    df["_sect"] = df[section].str.casefold()
    df["_lname"] = df[lname].str.casefold()

    df = df.sort_values(
        by=["_dept", "_subj", "_num", "_sect", "_lname"],
        kind="mergesort"
    )

    return df.drop(columns=["_dept", "_subj", "_num", "_sect", "_lname"])


# ---------------------------
# GUI
# ---------------------------

def get_inputs():
    root = tk.Tk()
    root.title("Canvas Report Setup")
    center_window(root, 400, 250)

    result = {"semester": None, "option": None}

    tk.Label(root, text="Semester:").pack()
    entry = tk.Entry(root)
    entry.pack()

    tk.Label(root, text="Report Type:").pack()

    option = tk.StringVar(value="Start of Term")

    tk.Radiobutton(root, text="Start of Term", value="Start of Term", variable=option).pack(anchor="w")
    tk.Radiobutton(root, text="Interim", value="Interim", variable=option).pack(anchor="w")
    tk.Radiobutton(root, text="End of Term", value="End of Term", variable=option).pack(anchor="w")

    def submit():
        if not entry.get().strip():
            messagebox.showerror("Error", "Enter semester")
            return
        result["semester"] = entry.get().strip()
        result["option"] = option.get()
        root.destroy()

    tk.Button(root, text="Submit", command=submit).pack(pady=10)

    root.mainloop()
    return result


inputs = get_inputs()
semester = inputs["semester"]
selected_option = inputs["option"]


# ---------------------------
# File picker
# ---------------------------

root = tk.Tk()
root.withdraw()

file = filedialog.askopenfilename(
    title="Select CSV",
    filetypes=[("CSV", "*.csv *.txt")]
)

if not file:
    raise SystemExit(0)


# ---------------------------
# Config
# ---------------------------

output_folder = os.path.join(os.path.dirname(file), "Canvas Presence Reports")
os.makedirs(output_folder, exist_ok=True)

canvas_col = "Canvas Presence"
dept_col = "academic_unit"
lname_col = "last_name"
subject_col = "course_subject"
number_col = "course_number"
section_col = "section_code"
college_col = "College"


# ---------------------------
# Load data
# ---------------------------

df = pd.read_csv(file, dtype=str)
df.columns = df.columns.str.strip()

df[canvas_col] = df.get("lms_section_id", "").apply(
    lambda x: "Yes" if pd.notna(x) and str(x).strip() else "No"
)

ensure_columns(df, [dept_col, subject_col, number_col, lname_col, section_col])


# ---------------------------
# College mapping (restore full version)
# ---------------------------

acad_dept_data = get_academic_department_data()

dept_to_college = {d: c for c, ds in acad_dept_data.items() for d in ds}
df[college_col] = df[dept_col].map(dept_to_college).fillna("???")


# ---------------------------
# SORT
# ---------------------------

df = excel_like_sort(df, dept_col, subject_col, number_col, lname_col, section_col)


# ---------------------------
# Split files
# ---------------------------

for college in df[college_col].unique():
    sub = df[df[college_col] == college].copy()

    path = os.path.join(output_folder, f"{college} {semester}.xlsx")

    wb = Workbook()

    # ---------------------------
    # Instructions sheet
    # ---------------------------
    ws1 = wb.active
    ws1.title = "Instructions"

    # Turn off gridlines
    ws1.sheet_view.showGridLines = False

    # Column widths
    ws1.column_dimensions["A"].width = 10
    ws1.column_dimensions["B"].width = 200

    instructions = get_instructions(college, semester, selected_option)

    # Font settings
    title_font = Font(
        name="Calibri (Body)",
        size=24,
        bold=True
    )
    subtitle_font = Font(
        name="Calibri (Body)",
        size=12,
        italic=True
    )
    h1_font = Font(
        name="Calibri (Body)",
        size=14,
        bold=True
    )
    h2_font = Font(
        name="Calibri (Body)",
        size=12,
        bold=True
    )
    text_font = Font(
        name="Calibri (Body)",
        size=12
    )

    row = 1

    for item in instructions:
        style, text, height = item

        ws1.cell(row=row, column=1, value=None)
        cell = ws1.cell(row=row, column=2, value=text)

        cell.alignment = Alignment(wrap_text=True, vertical="top")

        if style == "TITLE":
            cell.font = title_font
        elif style == "SUBTITLE":
            cell.font = subtitle_font
        elif style == "H1":
            cell.font = h1_font
        elif style == "H2":
            cell.font = h2_font
        elif style == "TEXT":
            cell.font = text_font
        elif style == "META":
            cell.font = Font(italic=True, size=10)

        # 👇 THIS is now fully controlled from instructions.py
        if height is not None:
            ws1.row_dimensions[row].height = height

        row += 1

    # ---------------------------
    # Data sheet
    # ---------------------------
    ws2 = wb.create_sheet("Canvas Presence Report")

    for c, col in enumerate(sub.columns, 1):
        ws2.cell(1, c, col)

    for r, row_data in enumerate(sub.itertuples(index=False), 2):
        for c, val in enumerate(row_data, 1):
            ws2.cell(r, c, val)

    max_row = ws2.max_row
    max_col = ws2.max_column

    canvas_idx = None
    for i, cell in enumerate(ws2[1], 1):
        if cell.value == canvas_col:
            canvas_idx = i

    if canvas_idx:
        apply_conditional_formatting(ws2, canvas_idx, max_row, max_col)

    create_table(ws2, max_row, max_col, f"{college}_Table")

    wb.save(path)


messagebox.showinfo("Done", f"Saved to:\n{output_folder}")